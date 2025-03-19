import openpyxl
import os

def save_answers_to_excel(pdf_filename, questions, responses):
    """
    Crea un archivo Excel con el nombre del PDF (misma ruta, distinta extensión .xlsx)
    y guarda las preguntas y respuestas en dos columnas: 'Pregunta' y 'Respuesta'.
    """
    # Crear un nuevo libro de Excel
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Respuestas Incidente"

    # Escribir encabezados en la primera fila
    sheet.cell(row=1, column=1, value="Pregunta")
    sheet.cell(row=1, column=2, value="Respuesta")

    # Rellenar con las preguntas y respuestas
    row_index = 2
    for i, question in enumerate(questions, start=1):
        sheet.cell(row=row_index, column=1, value=f"{i}. {question}")  # Pregunta
        if question in responses:
            sheet.cell(row=row_index, column=2, value=responses[question])  # Respuesta
        else:
            sheet.cell(row=row_index, column=2, value="No se encontró respuesta.")
        row_index += 1

    # Obtener el nombre base (sin extensión) del PDF
    excel_filename = os.path.splitext(pdf_filename)[0] + ".xlsx"

    # Guardar el Excel en disco
    workbook.save(excel_filename)
    return excel_filename
